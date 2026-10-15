import io
import json
import os
import tempfile
import unittest
import zipfile
from unittest.mock import patch

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from server.app import app, store
from server.dictionary import DICTIONARY


def pdf_bytes(text, encrypted=False):
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_textbox(fitz.Rect(40, 40, 550, 800), text, fontsize=11)
    kwargs = {"encryption": fitz.PDF_ENCRYPT_AES_256, "owner_pw": "owner", "user_pw": "secret"} if encrypted else {}
    data = document.tobytes(**kwargs)
    document.close()
    return data


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="tcc-api-")
        self.environment = patch.dict(os.environ, {"TCC_DATA_DIR": self.directory.name})
        self.environment.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.environment.stop()
        self.directory.cleanup()

    def demo(self):
        response = self.client.post("/api/demo")
        self.assertEqual(response.status_code, 201, response.text)
        batch = response.json()
        self.assertEqual(batch["status"], "completed", batch.get("error"))
        return batch

    def new_request(self):
        return {"config": {"source": "csv", "methods": ["M1", "M2", "A1"], "target": 2},
                "ppc": {"filename": "ppc.pdf", "pages": 1, "size": 100, "disciplines": [{"id": "d1", "name": "Programação", "period": 1, "hours": 60, "page": 1,
                                                                                         "syllabus": "Programação em Python. SQL e banco de dados.", "objectives": "Implementar algoritmos", "bibliography": "Docker e Kubernetes"}]},
                "jobs": [{"id": "ignored", "title": "Dev Python", "company": "Empresa", "city": "Itajaí", "description": "Python e Docker", "source": "csv"}]}

    def test_health_no_credentials_required(self):
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["dictionary_count"], 73)
        self.assertNotIn('"api_key":', response.text.lower())

    def test_pdf_text_segmentation_excludes_bibliography(self):
        text = "Disciplina: Programacao\nPeriodo: 1\nCarga horaria: 60\nEmenta: Logica de programacao em Python, variaveis e estruturas de dados.\nObjetivos: Desenvolver algoritmos e compreender estruturas de repeticao.\nBibliografia: Livro de Docker e Kubernetes."
        response = self.client.post("/api/ppc", files={"file": ("../../ppc.pdf", pdf_bytes(text), "application/pdf")})
        self.assertEqual(response.status_code, 200, response.text)
        ppc = response.json()
        self.assertEqual(ppc["filename"], "ppc.pdf")
        discipline = ppc["disciplines"][0]
        self.assertEqual(discipline["name"], "Programacao")
        self.assertIn("Python", discipline["syllabus"])
        self.assertNotIn("Docker", discipline["syllabus"])
        self.assertIn("Docker", discipline["bibliography"])

    def test_scanned_encrypted_and_invalid_pdf_refused(self):
        for data, expected in ((pdf_bytes(""), "OCR"), (pdf_bytes("Protected", encrypted=True), "senha"), (b"not a PDF", "PDF")):
            response = self.client.post("/api/ppc", files={"file": ("ppc.pdf", data, "application/pdf")})
            self.assertEqual(response.status_code, 422)
            self.assertIn(expected, response.json()["detail"])

    def test_jobs_import_dedup_and_brazilian_headers(self):
        data = "cargo;empresa;descricao;cidade\nDev;Acme;Python e SQL;Itajai\nDev;Acme;Python e SQL;Itajai\nSuporte;Outra;Linux e redes;Itajai\nSem descricao;Acme;;Itajai\n"
        response = self.client.post("/api/jobs/preview", files={"file": ("vagas.csv", data.encode(), "text/csv")})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 2)
        self.assertEqual(response.json()["duplicates"], 1)
        self.assertTrue(response.json()["warnings"])
        bad = self.client.post("/api/jobs/preview", files={"file": ("vagas.csv", b"name\nPython", "text/csv")})
        self.assertEqual(bad.status_code, 422)

    def test_full_review_confirmation_and_immutable_history(self):
        response = self.client.post("/api/batches", json=self.new_request())
        self.assertEqual(response.status_code, 201, response.text)
        batch = response.json()
        self.assertEqual(batch["status"], "review")
        batch_id = batch["id"]
        discipline = batch["disciplines"][0]
        discipline["syllabus"] += " Testes de software."
        edit = self.client.patch(f"/api/batches/{batch_id}/disciplines", json={"disciplines": [discipline], "corrections": 1})
        self.assertEqual(edit.status_code, 200)
        confirm = self.client.post(f"/api/batches/{batch_id}/confirm")
        self.assertEqual(confirm.status_code, 200)
        result = self.client.get(f"/api/batches/{batch_id}").json()
        self.assertEqual(result["status"], "completed", result.get("error"))
        self.assertEqual(len(result["stages"]), 6)
        self.assertEqual(next(item for item in result["stages"] if item["key"] == "classification")["status"], "completed")
        self.assertEqual(result["results"]["summary"]["jobs"], 1)
        self.assertEqual(result["validation"]["extraction"], [])
        self.assertIsNone(result["validation"]["agreement"])
        self.assertEqual(self.client.patch(f"/api/batches/{batch_id}/disciplines", json={"disciplines": [discipline]}).status_code, 409)
        self.assertEqual(self.client.post(f"/api/batches/{batch_id}/confirm").status_code, 409)
        association = result["results"]["associations"][0]
        decision = self.client.post(f"/api/batches/{batch_id}/decisions", json={"association_id": association["id"], "decision": "rejected"})
        self.assertEqual(decision.status_code, 200)
        self.assertEqual(decision.json()["results"]["associations"][0]["human_decision"], "rejected")
        self.assertEqual(decision.json()["results"]["indicators"], result["results"]["indicators"])
        self.assertEqual(len(self.client.get("/api/batches").json()), 1)

    def test_cancel_does_not_publish_results(self):
        batch = self.client.post("/api/batches", json=self.new_request()).json()
        response = self.client.post(f"/api/batches/{batch['id']}/cancel")
        self.assertEqual(response.json()["status"], "cancelled")
        self.assertIsNone(response.json()["results"])
        self.assertEqual(self.client.post(f"/api/batches/{batch['id']}/confirm").status_code, 409)
        self.assertEqual(self.client.get(f"/api/batches/{batch['id']}/export").status_code, 409)

    def test_exports_blind_sample_and_real_validation(self):
        batch = self.demo()
        batch_id = batch["id"]
        xlsx = self.client.get(f"/api/batches/{batch_id}/export?format=xlsx")
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(len(load_workbook(io.BytesIO(xlsx.content)).sheetnames), 6)
        archive = self.client.get(f"/api/batches/{batch_id}/export?format=csv&tables=jobs,indicators")
        self.assertEqual(set(zipfile.ZipFile(io.BytesIO(archive.content)).namelist()), {"jobs.csv", "indicators.csv"})
        sample = self.client.get(f"/api/batches/{batch_id}/sample")
        book = load_workbook(io.BytesIO(sample.content))
        self.assertEqual(set(book.sheetnames), {"extraction", "associations", "guide"})
        self.assertNotIn("prediction", json.dumps([row for row in book["extraction"].iter_rows(min_row=1, max_row=1, values_only=True)]))
        labels = book["extraction"]
        header = [cell.value for cell in labels[1]]
        label_column = header.index("present") + 1
        labels.cell(2, label_column, 1)
        labels.cell(3, label_column, 0)
        # A label beyond 2000 rows proves annotation files are not truncated like vacancy previews.
        labels.cell(2200, label_column, 1)
        association_headers = [cell.value for cell in book["associations"][1]]
        book["associations"].cell(2, association_headers.index("related") + 1, 1)
        output = io.BytesIO()
        book.save(output)
        primary = self.client.post(f"/api/batches/{batch_id}/annotations", files={"file": ("human.xlsx", output.getvalue())}, data={"annotator": "primary"})
        self.assertEqual(primary.status_code, 200, primary.text)
        validation = primary.json()["validation"]
        self.assertEqual(validation["status"], "annotated")
        self.assertEqual(validation["primary"]["counts"]["labels"], 4)
        self.assertTrue(validation["extraction"])
        self.assertTrue(validation["association"])
        self.assertIsNone(validation["agreement"])
        secondary = self.client.post(f"/api/batches/{batch_id}/annotations", files={"file": ("second.xlsx", output.getvalue())}, data={"annotator": "secondary"})
        self.assertEqual(secondary.status_code, 200)
        self.assertEqual(secondary.json()["validation"]["agreement"]["kappa"], 1)
        self.assertEqual(secondary.json()["validation"]["agreement"]["items"], 4)

    def test_dictionary_only_affects_future_batches(self):
        batch = self.demo()
        fingerprint = batch["results"]["fingerprint"]
        entry = DICTIONARY[0]
        response = self.client.post("/api/dictionary", json={"competency_id": entry["id"], "alias": "NovaAliasTCC"})
        self.assertEqual(response.status_code, 200)
        new_dictionary = self.client.get("/api/dictionary").json()["entries"]
        self.assertIn("NovaAliasTCC", next(item for item in new_dictionary if item["id"] == entry["id"])["synonyms"])
        self.assertEqual(self.client.get(f"/api/batches/{batch['id']}").json()["results"]["fingerprint"], fingerprint)
        self.assertEqual(self.client.post("/api/dictionary", json={"competency_id": "invalid", "alias": "x"}).status_code, 422)

    def test_csv_formula_injection_sanitized(self):
        request = self.new_request()
        request["jobs"][0]["title"] = '=HYPERLINK("https://example.invalid")'
        batch = self.client.post("/api/batches", json=request).json()
        self.client.post(f"/api/batches/{batch['id']}/confirm")
        response = self.client.get(f"/api/batches/{batch['id']}/export?format=xlsx&tables=jobs")
        book = load_workbook(io.BytesIO(response.content))
        row = list(book.active.values)[1]
        self.assertTrue(str(row[0]).startswith("'="))

    def test_input_validation_and_missing_batch(self):
        self.assertEqual(self.client.get("/api/batches/12345").status_code, 404)
        request = self.new_request()
        request["config"]["target"] = 999999
        self.assertEqual(self.client.post("/api/batches", json=request).status_code, 422)
        request = self.new_request()
        request["jobs"] = []
        self.assertEqual(self.client.post("/api/batches", json=request).status_code, 422)
        request = self.new_request()
        request["ppc"]["disciplines"][0]["syllabus"] = ""
        request["ppc"]["disciplines"][0]["objectives"] = ""
        self.assertEqual(self.client.post("/api/batches", json=request).status_code, 422)

    def test_sample_can_be_frozen_before_methods_and_edit_requires_new_sample(self):
        batch = self.client.post("/api/batches", json=self.new_request()).json()
        sample = self.client.get(f"/api/batches/{batch['id']}/sample")
        self.assertEqual(sample.status_code, 200)
        self.assertTrue(store().get(batch["id"]).get("sample"))
        self.assertIsNone(store().get(batch["id"])["results"])
        response = self.client.patch(f"/api/batches/{batch['id']}/disciplines", json={"disciplines": batch["disciplines"], "corrections": 1})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("sample", store().get(batch["id"]))

    def test_startup_marks_only_interrupted_work(self):
        completed = self.demo()
        review = self.client.post("/api/batches", json=self.new_request()).json()
        interrupted = self.client.post("/api/batches", json=self.new_request()).json()
        collecting = self.client.post("/api/batches", json=self.new_request()).json()
        store().update(interrupted["id"], lambda batch: batch.update(status="processing"))
        store().update(collecting["id"], lambda batch: batch["collection"].update(running=True))
        with TestClient(app) as restarted:
            self.assertEqual(restarted.get(f"/api/batches/{completed['id']}").json()["status"], "completed")
            self.assertEqual(restarted.get(f"/api/batches/{review['id']}").json()["status"], "review")
            failed = restarted.get(f"/api/batches/{interrupted['id']}").json()
            self.assertEqual(failed["status"], "failed")
            self.assertIn("reinício", failed["error"])
            partial = restarted.get(f"/api/batches/{collecting['id']}").json()
            self.assertEqual(partial["status"], "review")
            self.assertFalse(partial["collection"]["running"])
            self.assertTrue(partial["collection"]["partial"])
            self.assertEqual(len(partial["jobs"]), 1)

    def test_gap_export_respects_configured_demand_threshold(self):
        request = self.new_request()
        request["config"]["threshold"] = 60
        request["jobs"].append({"title": "Dev Python outro", "description": "Python e Linux"})
        batch = self.client.post("/api/batches", json=request).json()
        self.client.post(f"/api/batches/{batch['id']}/confirm")
        response = self.client.get(f"/api/batches/{batch['id']}/export?tables=gaps")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book["gaps"].max_row, 1)

    def test_explicit_area_and_title_priority_survive_processing(self):
        request = self.new_request()
        request["jobs"] = [{"title": "Desenvolvedor Java", "description": "Banco de dados e redes TCP/IP", "area": "Backend"}]
        batch = self.client.post("/api/batches", json=request).json()
        self.assertEqual(batch["jobs"][0]["area"], "Programação")
        self.client.post(f"/api/batches/{batch['id']}/confirm")
        completed = self.client.get(f"/api/batches/{batch['id']}").json()
        self.assertEqual(completed["jobs"][0]["area"], "Programação")


if __name__ == "__main__":
    unittest.main()
